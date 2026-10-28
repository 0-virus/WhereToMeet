"""공식 노선 XLSX와 기존 정류소·9월 관측 대조. 읽기 전용 openpyxl 필요."""
import hashlib
import json
import math
from collections import defaultdict
from pathlib import Path
import runpy
import sys
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
XLSX_HASH = 'b54816e9b7e77bcf49dd7562b993db06a093daa666351f115ea4d1dd0dd8b1df'
PRIOR = 'data/review-evidence/bus-connectivity-2026-10-29.json'
PRIOR_HASH = 'ae5cef26a6a6d8ba22678f441eda8658dc5e81b394a2e136ae0a6c2653c041f2'
distance = runpy.run_path(str(ROOT/'scripts/audit-bus-connectivity-20261029.py'))['distance']


def links(rows):
    grouped = defaultdict(list)
    for row in rows:
        grouped[row['routeId']].append(row)
    return [{'routeId': key, 'name': grouped[key][0]['routeName'],
             'sequences': sorted(r['sequence'] for r in grouped[key])} for key in sorted(grouped)]


def main(xlsx_path, output):
    blob = Path(xlsx_path).read_bytes()
    assert hashlib.sha256(blob).hexdigest() == XLSX_HASH, '다른 날짜/변경된 노선 원문'
    prior_text = (ROOT/PRIOR).read_text(encoding='utf-8')
    assert hashlib.sha256(prior_text.encode()).hexdigest() == PRIOR_HASH, '관측 근거 변경'
    prior = json.loads(prior_text)
    for source in prior['locationSources']:
        for file in source['files']:
            assert hashlib.sha256((ROOT/file['path']).read_bytes()).hexdigest() == file['sha256']
    raw = {r['NODE_ID']: r for r in json.loads((ROOT/'data/raw/seoul-bus-20260902.json').read_text(encoding='utf-8'))}
    workbook = load_workbook(xlsx_path, read_only=True, data_only=False)
    try:
        assert workbook.sheetnames == ['Data']
        sheet = workbook['Data']
        assert sheet.max_column == 8 and sheet.max_row == 41689
        iterator = sheet.iter_rows(values_only=True)
        assert next(iterator) == ('ROUTE_ID', '노선명', '순번', 'NODE_ID', 'ARS_ID', '정류소명', 'X좌표', 'Y좌표')
        rows = []
        for route, name, seq, stop, ars, stop_name, lng, lat in iterator:
            assert all(type(v) is int and 100000000 <= v <= 999999999 for v in (route, stop))
            assert type(seq) is int and seq > 0
            assert all(isinstance(v, str) and v and not v.startswith('=') for v in (name, ars, stop_name))
            assert len(ars) in (4, 5) and ars.isascii() and ars.isdigit()
            assert all(type(v) in (int, float) and math.isfinite(v) for v in (lng, lat))
            assert -180 <= lng <= 180 and -90 <= lat <= 90
            rows.append({'routeId': str(route), 'routeName': name, 'sequence': seq,
                         'sourceId': str(stop), 'arsId': ars, 'name': stop_name, 'lng': lng, 'lat': lat})
    finally:
        workbook.close()
    assert len(rows) == 41688
    assert len({(r['routeId'], r['sequence']) for r in rows}) == len(rows), '노선 내 순번 중복'
    by_stop, by_route = defaultdict(list), defaultdict(list)
    for row in rows:
        by_stop[row['sourceId']].append(row)
        by_route[row['routeId']].append(row)
    assert all(len({r['routeName'] for r in group}) == 1 for group in by_route.values())
    assert all(len({(r['arsId'], r['name'], r['lng'], r['lat']) for r in group}) == 1 for group in by_stop.values())
    matched = by_stop.keys() & raw.keys()
    assert all(by_stop[key][0]['arsId'] == raw[key]['ARS_ID'] for key in matched)
    assert all(by_stop[key][0]['name'] == raw[key]['name'] for key in matched)
    coordinates = []
    for key in sorted(matched):
        source, original = by_stop[key][0], raw[key]
        if (source['lat'], source['lng']) != (original['lat'], original['lng']):
            coordinates.append({'sourceId': key, 'name': original['name'],
                                'locationSourcePoint': {k: original[k] for k in ('lat', 'lng')},
                                'routeSourcePoint': {k: source[k] for k in ('lat', 'lng')},
                                'straightMetres': distance(original, source), 'releaseEligible': False})
    differences = {r['sourceId']: r for r in coordinates}
    targets = {r['sourceId']: r for group in ('currentPreview', 'outsideCentralLaneTop30', 'outsideCentralLaneAtLeast1kmTop30')
               for r in prior[group]}
    for key in prior['previewIdsWithoutObservation']:
        targets[key] = {'observedRouteNumbers': [], 'inPreview': True}
    reviewed = []
    for key, previous in sorted(targets.items()):
        grouped = links(by_stop.get(key, []))
        observed = sorted({r['number'] for r in previous['observedRouteNumbers']})
        names = sorted({r['name'] for r in grouped})
        reviewed.append({'sourceId': key, 'name': raw[key]['name'], 'inPreview': previous['inPreview'],
                         'routeSourcePresent': bool(grouped), 'routeIdCount': len(grouped), 'routeLinks': grouped,
                         'observedRouteNumberCount': len(previous['observedRouteNumbers']),
                         'namesOnlyInRouteSource': sorted(set(names)-set(observed)),
                         'numbersOnlyInObservation': sorted(set(observed)-set(names)),
                         'coordinateDifferenceMetres': differences[key]['straightMetres'] if key in differences else (0 if grouped else None),
                         'releaseEligible': False})
    summary = {'routeRows': len(rows), 'routeIds': len(by_route),
               'routeSourceStopIds': len(by_stop), 'matchedLocationIds': len(matched),
               'routeStopsWithoutLocation': len(by_stop.keys()-raw.keys()), 'locationStopsWithoutRoute': len(raw.keys()-matched),
               'arsConflicts': 0, 'nameConflicts': 0, 'withinSourceLocationConflicts': 0,
               'duplicateRouteSequenceKeys': 0,
               'fourDigitArsRows': sum(len(r['arsId']) == 4 for r in rows),
               'repeatedRouteStopRows': len(rows)-len({(r['routeId'], r['sourceId']) for r in rows}),
               'coordinateDifferences': len(coordinates),
               'coordinateDifferenceOverMetres': {str(n): sum(r['straightMetres'] > n for r in coordinates) for n in (1, 10, 100)},
               'reviewTargets': len(reviewed), 'targetsWithRouteSource': sum(r['routeSourcePresent'] for r in reviewed),
               'targetsWithDifferentNameSets': sum(bool(r['namesOnlyInRouteSource'] or r['numbersOnlyInObservation']) for r in reviewed),
               'targetsWithCoordinateDifferences': sum(r['coordinateDifferenceMetres'] is not None and r['coordinateDifferenceMetres'] > 0 for r in reviewed)}
    assert summary['routeSourceStopIds'] == len(matched) + summary['routeStopsWithoutLocation']
    assert len(targets) == 120 and sum(r['inPreview'] for r in reviewed) == 61
    result = {'reviewedAt': '2026-10-29',
              'source': {'id': 'OA-1095', 'sourceUrl': 'https://data.seoul.go.kr/dataList/OA-1095/S/1/datasetView.do',
                         'filename': '서울시버스노선별정류소정보(20260902).xlsx', 'date': '2026-09-02',
                         'bytes': len(blob), 'sha256': XLSX_HASH, 'license': '공공누리 1유형·서울특별시 출처표시',
                         'downloadUrl': 'https://datafile.seoul.go.kr/bigfile/iot/inf/nio_download.do?&useCache=false',
                         'downloadForm': {'infId': 'OA-1095', 'seq': '58', 'seqNo': '', 'infSeq': '2'}},
              'priorEvidence': {'path': PRIOR, 'normalizedLfSha256': PRIOR_HASH},
              'method': ['정류소 ID 정확한 일치에 ARS·이름 대조. 노선 ID를 중복 제거하고 반복 순번은 보존.',
                         '원문의 경기 정류소 2행은 4자리 ARS이며 앞에 0을 붙이거나 서울 ARS로 변환하지 않음.',
                         '9월 2일 정적 노선 원본과 9월 한 달 관측의 노선명/번호 문자열 차이를 그대로 보존. A/B·접미사를 자동 결합하지 않음.',
                         '노선명과 관측 번호 비교는 수단 코드를 결합하지 않은 문자열 진단이며 실제 노선 동일성 승인 아님.',
                         '좌표 차이는 원본 간 Haversine 직선거리. 1/10/100m는 설명용 구간이며 수정·채택 기준 아님.',
                         '순번만으로 상행/하행·현재 운행·약속 시각 가용성·공공 대기 공간을 확정하지 않음.',
                         '정적 자료의 노선 ID는 현재 시간표·경기/인천 전 노선·보행 접근성·출시 승인을 보장하지 않음.'],
              'summary': summary, 'coordinateDifferencesOver100m': sorted((r for r in coordinates if r['straightMetres'] > 100), key=lambda r: (-r['straightMetres'], r['sourceId'])),
              'fourDigitArsRows': [r for r in rows if len(r['arsId']) == 4],
              'reviewedStops': reviewed}
    Path(output).write_text(json.dumps(result, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    assert links([{'routeId': '1', 'routeName': '가', 'sequence': 9},
                  {'routeId': '1', 'routeName': '가', 'sequence': 1}]) == [{'routeId': '1', 'name': '가', 'sequences': [1, 9]}]
    assert links([]) == []
    assert len(links([{'routeId': '1', 'routeName': '가', 'sequence': 1},
                      {'routeId': '2', 'routeName': '가', 'sequence': 1}])) == 2
    if len(sys.argv) != 3:
        raise SystemExit('사용법: python scripts/audit-bus-route-links-20261029.py 원문.xlsx 결과.json')
    main(sys.argv[1], sys.argv[2])
